"""Lee el Excel (data/Donde Comer.xlsx) y genera los JSON de datos en build/.

Se puede correr a mano (python scripts/build.py) o lo corre automáticamente el
GitHub Action cada vez que se sube un Excel nuevo.
"""
import json
import os
import re
import unicodedata

import openpyxl

from coords import COORDS
from geocode import geocode_places

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DATA_XLSX = os.path.join(ROOT, "data", "Donde Comer.xlsx")
CACHE_PATH = os.path.join(ROOT, "data", "geocode_cache.json")
OUT_DIR = os.path.join(ROOT, "build")

# Coordenadas de emergencia si aparece un barrio que todavía no está en coords.py
# (Obelisco, centro de CABA) — mejor que el pin desaparezca en el medio del mapa
# y no en el medio del río.
DEFAULT_COORDS = (-34.6037, -58.3816)


def slugify(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    s = s.lower().strip()
    s = re.sub(r"[^a-z0-9]+", "-", s)
    s = re.sub(r"-+", "-", s).strip("-")
    return s or "lugar"


def load_rows():
    wb = openpyxl.load_workbook(DATA_XLSX, data_only=True)
    ws = wb["Hoja1"] if "Hoja1" in wb.sheetnames else wb.worksheets[0]
    data = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        if not r:
            continue
        row = (list(r) + [None] * 12)[:12]
        _, cat, subcat, nombre, localidad, barrio, video, direccion, linkmapa, geo, visitada, comentarios = row
        if not nombre:
            continue
        data.append({
            "categoria": cat, "subcategoria": subcat, "nombre": nombre, "localidad": localidad,
            "barrio": barrio, "video": video, "direccion": direccion, "link_mapa": linkmapa,
            "geo": geo, "comentarios": comentarios,
        })
    return data


def build():
    data = load_rows()

    places = []
    seen_ids = {}
    for d in data:
        barrio = str(d["barrio"]).strip() if d.get("barrio") else "Sin barrio"
        fallback_lat, fallback_lon = COORDS.get(barrio, DEFAULT_COORDS)

        base_id = slugify(f"{d['nombre']}-{barrio}")
        if base_id in seen_ids:
            seen_ids[base_id] += 1
            pid = f"{base_id}-{seen_ids[base_id]}"
        else:
            seen_ids[base_id] = 1
            pid = base_id

        direccion = str(d["direccion"]) if d.get("direccion") else ""
        link_mapa = str(d["link_mapa"]) if d.get("link_mapa") else ""
        geo = str(d["geo"]) if d.get("geo") else ""
        localidad = str(d["localidad"]).strip() if d.get("localidad") else ""
        if not direccion and not geo:
            # lugares tipo "Varios (ver video)" sin dirección propia
            direccion = "Ver el video para la dirección"
        elif not geo and direccion:
            geo = f"{direccion}, {barrio}, {localidad or 'CABA'}, Argentina"

        places.append({
            "id": pid,
            "cat": str(d["categoria"]) if d.get("categoria") else "",
            "subcat": str(d["subcategoria"]) if d.get("subcategoria") else "",
            "nombre": str(d["nombre"]),
            "localidad": localidad,
            "barrio": barrio,
            "ig": str(d["video"]) if d.get("video") else "",
            "dir": direccion,
            "maps": link_mapa,
            "geo": geo,
            "tip": str(d["comentarios"]).strip() if d.get("comentarios") else "",
            "lat": fallback_lat,
            "lon": fallback_lon,
            "approx": True,
        })

    geocode_places(places, CACHE_PATH)

    cats = sorted(set(p["cat"] for p in places if p["cat"]))
    subcats = sorted(set(p["subcat"] for p in places if p["subcat"]))
    barrios = sorted(set(p["barrio"] for p in places))
    localidades = sorted(set(p["localidad"] for p in places if p["localidad"]))

    os.makedirs(OUT_DIR, exist_ok=True)
    for name, obj in [
        ("places", places), ("cats", cats), ("subcats", subcats),
        ("barrios", barrios), ("localidades", localidades),
    ]:
        with open(os.path.join(OUT_DIR, f"{name}.json"), "w", encoding="utf-8") as f:
            json.dump(obj, f, ensure_ascii=False)

    print(json.dumps({
        "n": len(places), "cats_n": len(cats), "subcats_n": len(subcats),
        "barrios_n": len(barrios), "localidades": localidades,
    }, ensure_ascii=False))


if __name__ == "__main__":
    build()
